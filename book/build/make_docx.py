# -*- coding: utf-8 -*-
"""
Build an editable RTL Word (.docx) version of the manual from index.html.
Captures headings, paragraphs, lists, tables, callout boxes, exercise cards,
figures (rasterized) and QR images. Text stays fully editable in Word.

Run after build.py:  python3 build/make_docx.py
"""
import os, sys, io, re, hashlib
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
BOOK = os.path.dirname(HERE)
sys.path.insert(0, HERE)
CACHE = os.path.join(HERE, "_svgcache"); os.makedirs(CACHE, exist_ok=True)

NAVY = RGBColor(0x0e, 0x2a, 0x47)
ACC  = RGBColor(0xc8, 0x43, 0x1a)
STEEL= RGBColor(0x5b, 0x6b, 0x7b)

# ---------- low-level RTL helpers ----------------------------------------
def _set_rtl_par(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:bidi'); pPr.append(b)

def _shade(p, color):
    pPr = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), color); pPr.append(sh)

def _run(p, text, bold=False, rtl=True, color=None, size=None, italic=False):
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.name = 'Arial'
    rpr = r._r.get_or_add_rPr()
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), 'Arial'); rf.set(qn('w:hAnsi'), 'Arial'); rf.set(qn('w:cs'), 'Arial')
    rpr.append(rf)
    if rtl:
        el = OxmlElement('w:rtl'); el.set(qn('w:val'), '1'); rpr.append(el)
    if color is not None: r.font.color.rgb = color
    if size is not None: r.font.size = Pt(size)
    return r

def _para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    _set_rtl_par(p)
    return p

# ---------- inline text with bold/en handling ----------------------------
def add_inline(p, node):
    if isinstance(node, NavigableString):
        t = str(node)
        if t.strip() or t == ' ':
            _run(p, t)
        return
    for child in node.children:
        if isinstance(child, NavigableString):
            t = str(child)
            if t.strip() or t == ' ':
                _run(p, t)
        elif isinstance(child, Tag):
            cls = child.get('class', [])
            if child.name in ('b', 'strong'):
                _run(p, child.get_text(), bold=True)
            elif child.name == 'em':
                _run(p, child.get_text(), color=ACC)
            elif child.name == 'u':
                r = _run(p, child.get_text()); r.underline = True
            elif child.name == 'sup':
                r = _run(p, child.get_text(), size=7)
                r.font.superscript = True
            elif 'en' in cls or 'ltr' in cls:
                _run(p, child.get_text(), rtl=False, bold=('en' in cls))
            elif child.name == 'br':
                _run(p, '\n')
            else:
                add_inline(p, child)

# ---------- SVG rasterization --------------------------------------------
def svg_to_png(svg_markup):
    key = hashlib.md5(svg_markup.encode('utf-8')).hexdigest()[:16]
    path = os.path.join(CACHE, key + '.png')
    if os.path.exists(path):
        return path
    from weasyprint import HTML
    import fitz
    m = re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg_markup)
    w, h = (float(m.group(1)), float(m.group(2))) if m else (300, 300)
    html = (f'<html><head><style>@page{{size:{w}px {h}px;margin:0}}'
            f'body{{margin:0}}</style></head><body>{svg_markup}</body></html>')
    pdf = HTML(string=html, base_url=BOOK + '/').write_pdf()
    doc = fitz.open(stream=pdf, filetype='pdf')
    doc[0].get_pixmap(dpi=150).save(path)
    return path

# ---------- element processors -------------------------------------------
def proc_table(doc, tbl):
    # caption
    cap = tbl.find('caption')
    if cap:
        p = _para(doc); _run(p, cap.get_text(), bold=True, color=NAVY, size=10)
    headers = [th.get_text() for th in tbl.select('thead th')]
    body_rows = tbl.select('tbody tr')
    if not headers:  # some tables without thead
        return
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = 2  # right
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        para = hdr[i].paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_rtl_par(para)
        _run(para, htext, bold=True, color=RGBColor(0xff,0xff,0xff), size=9)
        _shade(para, '0e2a47')
    for tr in body_rows:
        cells = tr.find_all('td')
        row = t.add_row().cells
        for i, td in enumerate(cells):
            if i >= len(headers): break
            para = row[i].paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_rtl_par(para)
            add_inline(para, td)
    doc.add_paragraph()

def proc_box(doc, box):
    title_el = box.find(class_='box-t')
    title = title_el.get_text() if title_el else ''
    p = _para(doc)
    if title:
        _run(p, '▍ ' + title, bold=True, color=ACC, size=10.5)
    _shade(p, 'f4f6f9')
    for child in box.children:
        if isinstance(child, Tag) and 'box-t' in child.get('class', []):
            continue
        if isinstance(child, Tag) and child.name in ('p',):
            pp = _para(doc); add_inline(pp, child); _shade(pp, 'f4f6f9')
        elif isinstance(child, Tag) and child.name in ('ul', 'ol'):
            proc_list(doc, child, shade='f4f6f9')
        elif isinstance(child, Tag) and child.name == 'small':
            pp = _para(doc); _run(pp, child.get_text(), color=STEEL, size=8.5); _shade(pp, 'f4f6f9')

def proc_list(doc, lst, shade=None):
    ordered = lst.name == 'ol'
    is_check = 'check' in lst.get('class', [])
    for i, li in enumerate(lst.find_all('li', recursive=False), 1):
        p = _para(doc)
        prefix = '☐  ' if is_check else (f'{i}.  ' if ordered else '•  ')
        _run(p, prefix, bold=not is_check, color=ACC)
        add_inline(p, li)
        if shade: _shade(p, shade)

def proc_exercise(doc, ex, compact=False):
    head = ex.find(class_='ex-head' if not compact else 'mx-head')
    name_el = ex.find(class_='exname')
    num_el = ex.find(class_='exn')
    p = _para(doc)
    if num_el: _run(p, num_el.get_text() + '  ', bold=True, color=RGBColor(0xff,0xff,0xff))
    _run(p, name_el.get_text() if name_el else 'תרגיל', bold=True,
         color=RGBColor(0xff,0xff,0xff), size=12)
    _shade(p, '163a5f')
    # dose chips
    chips = ex.select('.dose .chip') or ([ex.find(class_='mx-dose')] if ex.find(class_='mx-dose') else [])
    if chips:
        dp = _para(doc)
        _run(dp, ' · '.join(c.get_text(' ', strip=True) for c in chips if c),
             bold=True, color=ACC, size=9)
    # fields
    for f in ex.select('.ex-field, .mx-f'):
        lbl = f.find(class_='lbl')
        fp = _para(doc)
        if lbl:
            _run(fp, lbl.get_text() + ': ', bold=True, color=ACC, size=9.5)
        # value = field minus label
        for child in f.children:
            if isinstance(child, Tag) and 'lbl' in child.get('class', []):
                continue
            if isinstance(child, Tag) and child.name in ('ul', 'ol'):
                # inline the list items compactly
                items = [li.get_text(' ', strip=True) for li in child.find_all('li')]
                _run(fp, ' · '.join(items), size=9.5)
            else:
                if isinstance(child, NavigableString):
                    if child.strip(): _run(fp, str(child), size=9.5)
                elif isinstance(child, Tag):
                    add_inline(fp, child)
    # progression / regression
    for pr in ex.select('.prog .p, .prog .r, .mx-pr .p, .mx-pr .r'):
        pp = _para(doc); _run(pp, pr.get_text(' ', strip=True), size=9, color=STEEL)
    # figures + media inside exercise
    for fig in ex.select('svg'):
        _insert_svg(doc, fig)
    proc_media(doc, ex)

def _insert_svg(doc, svg):
    try:
        png = svg_to_png(str(svg))
        m = re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', str(svg))
        w = float(m.group(1)) if m else 260
        width_in = min(3.2, w / 96.0)
        p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run().add_picture(png, width=Inches(width_in))
    except Exception as e:
        pass

def proc_media(doc, container):
    for qb in container.select('.qrbox'):
        img = qb.find('img')
        qt = qb.find(class_='qt')
        src = qb.find(class_='src')
        lnk = qb.find(class_='lnk')
        tag = qb.find(class_='tag')
        p = _para(doc)
        label = (tag.get_text() + ' — ') if tag else ''
        _run(p, '▶ ' + label + (qt.get_text() if qt else ''), bold=True, color=NAVY, size=9.5)
        if src:
            _run(p, '  · מקור: ' + src.get_text(), color=ACC, size=8.5)
        if lnk:
            lp = _para(doc); _run(lp, lnk.get_text(), rtl=False, color=RGBColor(0x2b,0x5f,0x8a), size=8)
        if img and img.get('src'):
            ip = os.path.join(BOOK, img['src'])
            if os.path.exists(ip):
                pp = _para(doc)
                try: pp.add_run().add_picture(ip, width=Inches(0.85))
                except Exception: pass

def proc_refs(doc, refs):
    h = refs.find(['h3'])
    hp = _para(doc, style='Heading 2')
    _run(hp, h.get_text() if h else 'מקורות', bold=True, color=NAVY, size=12)
    for i, li in enumerate(refs.select('ol > li'), 1):
        p = _para(doc)
        _run(p, f'{i}. ', bold=True, color=ACC)
        for child in li.children:
            if isinstance(child, NavigableString):
                if child.strip(): _run(p, str(child), size=8.5)
            elif isinstance(child, Tag):
                if child.name == 'a':
                    _run(p, child.get_text(), rtl=False, color=RGBColor(0x2b,0x5f,0x8a), size=8)
                elif child.name == 'br':
                    pass
                else:
                    add_inline(p, child)

# ---------- section walker -----------------------------------------------
def walk(doc, node):
    for child in node.children:
        if not isinstance(child, Tag):
            continue
        cls = child.get('class', [])
        name = child.name
        if name == 'h2':
            if 'ctitle' in cls:
                # chapter opener handled by parent; still emit
                pass
            p = _para(doc, style='Heading 1')
            _run(p, child.get_text(), bold=True, color=NAVY, size=16)
        elif name == 'h3':
            p = _para(doc, style='Heading 2')
            _run(p, child.get_text().lstrip('▍ '), bold=True, color=RGBColor(0x16,0x3a,0x5f), size=12.5)
        elif name == 'h4':
            p = _para(doc, style='Heading 3')
            _run(p, child.get_text(), bold=True, color=ACC, size=11)
        elif name == 'p':
            p = _para(doc); add_inline(p, child)
        elif name in ('ul', 'ol'):
            proc_list(doc, child)
        elif name == 'table':
            proc_table(doc, child)
        elif name == 'figure':
            for svg in child.select('svg'):
                _insert_svg(doc, svg)
            fc = child.find('figcaption')
            if fc:
                p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
                _run(p, fc.get_text(), color=STEEL, size=8.5, italic=True)
        elif 'box' in cls:
            proc_box(doc, child)
        elif 'ex' in cls:
            proc_exercise(doc, child, compact=False)
        elif 'mx' in cls:
            proc_exercise(doc, child, compact=True)
        elif 'media' in cls:
            proc_media(doc, child)
        elif 'refs' in cls:
            proc_refs(doc, child)
        elif 'chap-opener' in cls:
            proc_chapter_opener(doc, child)
        elif 'chap-map' in cls:
            p = _para(doc); _run(p, child.get_text(' ', strip=True), color=STEEL, size=9)
        elif name in ('section', 'div'):
            walk(doc, child)

def proc_chapter_opener(doc, opener):
    doc.add_page_break()
    cnum = opener.find(class_='cnum')
    ctitle = opener.find(class_='ctitle')
    lead = opener.find(class_='lead')
    if cnum:
        p = _para(doc); _run(p, cnum.get_text(), bold=True, color=ACC, size=11)
    if ctitle:
        p = _para(doc, style='Heading 1'); _run(p, ctitle.get_text(), bold=True, color=NAVY, size=20)
    if lead:
        p = _para(doc); _run(p, lead.get_text(), italic=True, color=RGBColor(0x16,0x3a,0x5f), size=11.5)

def proc_cover(doc, cover):
    for _ in range(3): doc.add_paragraph()
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, 'יסודות התנועה בכדורסל', bold=True, color=NAVY, size=30)
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, 'Movement Foundation for Basketball', rtl=False, bold=True, color=ACC, size=15)
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, 'שיפור תנועה · בלימה · הורדת מרכז כובד · שינויי כיוון · מניעת פציעות',
         color=STEEL, size=12)
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, 'מדריך עבודה מקצועי למאמן ביצועים וספורטתרפיסט · גילאי 13–20', color=NAVY, size=11)

def main():
    with open(os.path.join(BOOK, 'index.html'), encoding='utf-8') as f:
        soup = BeautifulSoup(f.read(), 'lxml')
    doc = Document()
    # base style RTL-friendly
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'; normal.font.size = Pt(10.5)
    body = soup.body
    for section in body.find_all('section', recursive=False):
        pass
    # process cover first
    cover = body.find(class_='cover')
    if cover: proc_cover(doc, cover)
    for child in body.children:
        if not isinstance(child, Tag): continue
        if 'cover' in child.get('class', []): continue
        walk(doc, child)
    out = os.path.join(BOOK, 'Movement_Foundation_Basketball.docx')
    doc.save(out)
    print('wrote', out, f'({os.path.getsize(out)//1024} KB)')

if __name__ == '__main__':
    main()
